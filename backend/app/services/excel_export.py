from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_excel(meeting: dict, agenda_items: list[dict], speaker_tasks: list[dict], validations: list[dict]) -> BytesIO:
    wb = Workbook()
    schedule = wb.active
    schedule.title = "正式日程表"
    guests = wb.create_sheet("嘉宾任务表")
    notes = wb.create_sheet("策划说明表")

    write_title(schedule, meeting["name"])
    schedule.append(["时间", "会议板块", "环节形式", "讲题", "主持人", "讲者", "讨论嘉宾", "讨论议题", "会议地点", "执行备注"])
    for item in agenda_items:
        schedule.append([
            f"{item['start_time']}-{item['end_time']}",
            item["section_name"],
            item["session_format"],
            item["topic"],
            item.get("host", ""),
            item.get("speaker", ""),
            item.get("panelists", ""),
            item.get("discussion_topic", ""),
            item["location"],
            item.get("execution_notes", ""),
        ])

    write_title(guests, "嘉宾任务表")
    guests.append(["嘉宾", "角色", "任务", "时间", "备注"])
    live_tasks = derive_speaker_tasks(agenda_items) or speaker_tasks
    for task in live_tasks:
        guests.append([task["expert_name"], task["role"], task["task_description"], task["time_range"], task.get("notes", "")])

    write_title(notes, "策划说明表")
    notes.append(["字段", "内容"])
    notes.append(["会议主题", meeting.get("theme", "")])
    notes.append(["日期", meeting["date"]])
    notes.append(["地点", meeting["location"]])
    notes.append(["会议级别", meeting["level"]])
    notes.append(["会议形式", meeting["format"]])
    notes.append(["学术领域", meeting["academic_field"]])
    notes.append(["会议类型", meeting["meeting_type"]])
    notes.append(["会议目标", meeting["goal"]])
    notes.append(["参会人群", meeting["audience"]])
    notes.append(["人数", meeting["attendee_count"]])
    notes.append(["策划说明", meeting.get("planning_notes", "")])
    notes.append([])
    notes.append(["校验等级", "校验信息"])
    for item in validations:
        notes.append([item["severity"], item["message"]])

    for ws in wb.worksheets:
        style_sheet(ws)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def derive_speaker_tasks(agenda_items: list[dict]) -> list[dict]:
    tasks: list[dict] = []
    role_fields = [("主持", "host"), ("讲者", "speaker"), ("讨论嘉宾", "panelists")]
    for item in agenda_items:
        time_range = f"{item['start_time']}-{item['end_time']}"
        for role, field in role_fields:
            raw_names = str(item.get(field, "") or "")
            names = [name.strip() for name in raw_names.replace("、", ",").split(",") if name.strip()]
            for name in names:
                tasks.append(
                    {
                        "expert_name": name,
                        "role": role,
                        "task_description": item.get("topic", ""),
                        "time_range": time_range,
                        "notes": item.get("execution_notes", ""),
                    }
                )
    return tasks


def write_title(ws, title: str) -> None:
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="176B87")
    ws["A1"].alignment = Alignment(horizontal="center")


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="EAF2F5")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 2:
                cell.font = Font(bold=True)
                cell.fill = header_fill
    widths = [16, 16, 14, 28, 16, 16, 18, 34, 20, 24]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
