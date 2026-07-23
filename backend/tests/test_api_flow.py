from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest

os.environ["MEETING_DB_PATH"] = str(Path(tempfile.gettempdir()) / "meeting_planner_test.db")

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


class ApiFlowTests(unittest.TestCase):
    def test_create_generate_validate_export(self) -> None:
        payload = {
            "name": "肺癌精准诊疗城市会",
            "date": "2026-08-06",
            "start_time": "14:00",
            "end_time": "17:00",
            "location": "会议酒店",
            "level": "区域级",
            "format": "线下",
            "academic_field": "肺癌",
            "meeting_type": "城市会",
            "goal": "提升围术期精准诊疗交流质量",
            "audience": "肿瘤科、胸外科专家",
            "attendee_count": 50,
        }

        with TestClient(app) as client:
            created = client.post("/api/meetings", json=payload)
            self.assertEqual(created.status_code, 200)
            meeting_id = created.json()["id"]

            generated = client.post(f"/api/meetings/{meeting_id}/generate")
            self.assertEqual(generated.status_code, 200)
            self.assertGreaterEqual(len(generated.json()["agenda_items"]), 5)

            validated = client.post(f"/api/meetings/{meeting_id}/validate")
            self.assertEqual(validated.status_code, 200)
            self.assertIn("validation_results", validated.json())

            exported = client.get(f"/api/meetings/{meeting_id}/export.xlsx")
            self.assertEqual(exported.status_code, 200)
            workbook = load_workbook(BytesIO(exported.content), read_only=True)
            self.assertEqual(workbook.sheetnames, ["正式日程表", "嘉宾任务表", "策划说明表"])


if __name__ == "__main__":
    unittest.main()
